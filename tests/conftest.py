import csv
import hashlib
import io
import json
import os
import re
import sqlite3
from collections.abc import Callable, Iterator
from datetime import date
from pathlib import Path

import httpx
import pytest
from pytest_httpx import HTTPXMock

from orrery import db, runs
from orrery.config import Settings
from orrery.ingest.awards import AwardsResult
from orrery.ingest.bulk import COLUMNS
from orrery.ingest.entities import EntitiesResult
from orrery.ingest.notices import ingest_notices
from orrery.sam.client import SamClient

# The Anthropic SDK talks httpx2, which pytest-httpx does not intercept: a stray real call
# must fail fast rather than reach the network.
os.environ.setdefault("ANTHROPIC_BASE_URL", "http://127.0.0.1:9")
os.environ.pop("ANTHROPIC_API_KEY", None)
for _name in [n for n in os.environ if n.startswith("ORRERY_")]:
    del os.environ[_name]  # a developer's shell must not leak settings into the tests

SEARCH_FIXTURE = json.loads(
    (Path(__file__).with_name("fixtures") / "sam_search_v2.json").read_text()
)
SEARCH_URL = re.compile(r".*/opportunities/v2/search.*")


NOW = "2026-09-09T00:00:00Z"
"""The instant every test runs at.

The notice fixtures carry absolute response deadlines, so without this the suite's result
depends on when it is run: on 2026-09-09 at 16:00 UTC one fixture deadline passed and four
tests went red with no code change behind them. Freezing the clock also makes stored
timestamps deterministic. A test that needs a different instant overrides db.utcnow itself,
as several already do."""


@pytest.fixture(autouse=True)
def frozen_clock(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(db, "utcnow", lambda: NOW)


@pytest.fixture
def db_path(tmp_path: Path) -> Path:
    return tmp_path / "orrery.sqlite"


@pytest.fixture
def conn(db_path: Path) -> Iterator[sqlite3.Connection]:
    """A migrated connection to a fresh database file."""
    connection = db.connect(db_path)
    db.migrate(connection)
    yield connection
    connection.close()


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        _env_file=None,
        data_dir=tmp_path,
        sam_api_key="test-key",
        sam_daily_budget=10,
        naics=["541512"],
        fetch_delay=0.0,
    )


@pytest.fixture
def run_id(conn: sqlite3.Connection) -> int:
    return runs.start(conn)


@pytest.fixture
def client(settings: Settings, conn: sqlite3.Connection, run_id: int) -> Iterator[SamClient]:
    with SamClient(settings, conn, run_id) as sam:
        yield sam


@pytest.fixture
def seed(
    conn: sqlite3.Connection, settings: Settings, httpx_mock: HTTPXMock
) -> Callable[[dict | None], None]:
    """Ingest the search fixture (or a modified copy); spends 1 of the 10-request budget."""

    def _seed(payload: dict | None = None) -> None:
        httpx_mock.add_response(url=SEARCH_URL, json=payload or SEARCH_FIXTURE)
        ingest_notices(conn, settings, posted_from=date(2026, 9, 5), posted_to=date(2026, 9, 6))

    return _seed


def _make_pdf(pages: list[str]) -> bytes:
    """A minimal PDF 1.4 (Helvetica, one text run per page) that pypdf extracts verbatim."""
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
            f" /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objs.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, obj in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


@pytest.fixture
def make_pdf() -> Callable[[list[str]], bytes]:
    return _make_pdf


def _make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """A .docx with the paragraphs, and a table after the first one when given, so document
    order and table flattening can both be asserted."""
    from docx import Document

    document = Document()
    for index, text in enumerate(paragraphs):
        document.add_paragraph(text)
        if table and index == 0:
            grid = document.add_table(rows=len(table), cols=len(table[0]))
            for row, cells in zip(grid.rows, table, strict=True):
                for cell, value in zip(row.cells, cells, strict=True):
                    cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def make_docx() -> Callable[..., bytes]:
    return _make_docx


def _make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    """A workbook of {sheet name: rows of cell values}, in the order given."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def make_xlsx() -> Callable[..., bytes]:
    return _make_xlsx


MOST_LINKS_NOTICE = max(
    SEARCH_FIXTURE["opportunitiesData"], key=lambda r: len(r["resourceLinks"] or [])
)["noticeId"]


@pytest.fixture
def fetched(
    conn: sqlite3.Connection, settings: Settings, seed: Callable[[dict | None], None]
) -> Callable[..., int]:
    """Seed the fixture notices, then store bytes as a fetched attachment; returns its id."""
    seed()

    def _fetched(filename: str, data: bytes, *, notice_id: str = MOST_LINKS_NOTICE) -> int:
        (attachment_id,) = conn.execute(
            "SELECT attachment_id FROM attachments WHERE notice_id = ? AND fetch_status = 'pending'"
            " ORDER BY attachment_id LIMIT 1",
            (notice_id,),
        ).fetchone()
        sha = hashlib.sha256(data).hexdigest()
        relative = Path("attachments") / notice_id / f"{sha[:16]}-{filename}"
        (settings.data_dir / relative).parent.mkdir(parents=True, exist_ok=True)
        (settings.data_dir / relative).write_bytes(data)
        conn.execute(
            "UPDATE attachments SET filename = ?, path = ?, content_hash = ?, fetched_at = ?,"
            " fetch_status = 'fetched' WHERE attachment_id = ?",
            (filename, relative.as_posix(), sha, db.utcnow(), attachment_id),
        )
        return attachment_id

    return _fetched


EMBED_URL = "http://localhost:11434/v1/embeddings"
FAKE_TERMS = ("xylophone", "zeppelin", "quokka")


def fake_vector(text: str) -> list[float]:
    """A 4-dim stand-in for a real embedding: term counts plus a constant, so cosine
    distance is meaningful and deterministic in tests."""
    lowered = text.lower()
    return [float(lowered.count(term)) for term in FAKE_TERMS] + [1.0]


def register_fake_embeddings(httpx_mock: HTTPXMock, batches: list[list[str]]) -> None:
    """Answer every embeddings request with fake vectors and record the batches seen."""

    def respond(request: httpx.Request) -> httpx.Response:
        inputs = json.loads(request.read())["input"]
        batches.append(list(inputs))
        data = [{"index": i, "embedding": fake_vector(t)} for i, t in enumerate(inputs)]
        return httpx.Response(200, json={"data": data})

    httpx_mock.add_callback(respond, url=EMBED_URL, is_reusable=True)


CHAT_URL = "http://localhost:11434/v1/chat/completions"


def register_fake_chat(httpx_mock: HTTPXMock, replies: list[str], requests: list[dict]) -> None:
    """Answer chat requests with canned replies in order (the last one repeats) and record
    every request body."""

    def respond(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.read())
        requests.append(body)
        text = replies[min(len(requests), len(replies)) - 1]
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"message": {"role": "assistant", "content": text}, "finish_reason": "stop"}
                ],
                "usage": {"prompt_tokens": 120, "completion_tokens": 40},
            },
        )

    httpx_mock.add_callback(respond, url=CHAT_URL, is_reusable=True)


@pytest.fixture
def fake_embeddings(httpx_mock: HTTPXMock) -> list[list[str]]:
    batches: list[list[str]] = []
    register_fake_embeddings(httpx_mock, batches)
    return batches


EXTRACT_DEFAULTS = {
    "Title": "Custodial Service",
    "Sol#": "75R60226Q00001",
    "Department/Ind.Agency": "HEALTH AND HUMAN SERVICES, DEPARTMENT OF",
    "CGAC": "075",
    "Sub-Tier": "HEALTH RESOURCES AND SERVICES ADMINISTRATION",
    "FPDS Code": "7526",
    "Office": "HRSA HEADQUARTERS",
    "AAC Code": "75R602",
    "PostedDate": "2026-09-05",
    "Type": "Solicitation",
    "BaseType": "Solicitation",
    "ArchiveType": "auto15",
    "ArchiveDate": "2026-09-30",
    "SetASideCode": "SBA",
    "SetASide": "Small Business Set Aside - Total",
    "ResponseDeadLine": "2026-09-15T15:00:00-04:00",
    "NaicsCode": "541512",
    "ClassificationCode": "S201",
    "PopState": "CO",
    "PopZip": "80503",
    "PopCountry": "USA",
    "Active": "Yes",
    "PrimaryContactFullname": "Point of Contact 1",
    "PrimaryContactEmail": "poc1@example.gov",
    "OrganizationType": "OFFICE",
    "State": "MD",
    "City": "ROCKVILLE",
    "ZipCode": "20852",
    "CountryCode": "USA",
    "Link": "https://sam.gov/workspace/contract/opp/x/view",
    "Description": "Section L \u2013 instructions\nline two",
}
_extract_counter = iter(range(1, 10_000))


def make_extract(rows: list[dict], columns: list[str] = COLUMNS) -> bytes:
    """A cp1252 extract with the real 47 headers; unspecified cells take defaults."""
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        full = {column: "" for column in columns}
        full.update({k: v for k, v in EXTRACT_DEFAULTS.items() if k in full})
        full.setdefault("NoticeId", "")
        if "NoticeId" in full and not full["NoticeId"]:
            full["NoticeId"] = f"{next(_extract_counter):032x}"
        full.update(row)
        writer.writerow(full)
    return out.getvalue().encode("cp1252")


AWARD_COLUMNS = next(
    csv.reader(
        io.StringIO(
            (Path(__file__).with_name("fixtures") / "usaspending_awards_sample.csv")
            .read_text(encoding="utf-8-sig")
            .splitlines()[0]
        )
    )
)
AWARD_DEFAULTS = {
    "award_id_piid": "75R60225F00001",
    "parent_award_id_piid": "47QTCA18D00L8",
    "current_total_value_of_award": "125000.50",
    "potential_total_value_of_award": "250000.00",
    "award_base_action_date": "2025-03-01",
    "award_latest_action_date": "2025-06-15",
    "period_of_performance_start_date": "2025-03-01",
    "period_of_performance_current_end_date": "2026-02-28",
    "period_of_performance_potential_end_date": "2027-02-28 00:00:00",
    "awarding_agency_code": "075",
    "awarding_agency_name": "Department of Health and Human Services",
    "awarding_sub_agency_code": "7526",
    "awarding_sub_agency_name": "Health Resources and Services Administration",
    "awarding_office_code": "75R602",
    "awarding_office_name": "HRSA HEADQUARTERS",
    "recipient_uei": "UE9QJD4KK1L6",
    "recipient_name": "LEIDOS, INC.",
    "cage_code": "5UTE1",
    "solicitation_identifier": "75R60225R00001",
    "naics_code": "541512",
    "product_or_service_code": "D399",
    "award_type_code": "C",
    "type_of_set_aside_code": "SBA",
    "extent_competed_code": "A",
    "usaspending_permalink": "https://www.usaspending.gov/award/CONT_AWD_EXAMPLE/",
    # Present in the real file and always dropped by the adapter.
    "recipient_phone_number": "5555550100",
    "highly_compensated_officer_1_name": "Officer Placeholder",
    "highly_compensated_officer_1_amount": "1",
}
_award_counter = iter(range(1, 10_000))


def make_awards_csv(rows: list[dict]) -> bytes:
    """A UTF-8 award summary with the real 286 headers; unspecified cells take defaults."""
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=AWARD_COLUMNS, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        full = dict.fromkeys(AWARD_COLUMNS, "")
        full.update(AWARD_DEFAULTS)
        full.update(row)
        if not full["contract_award_unique_key"]:
            full["contract_award_unique_key"] = f"CONT_AWD_{next(_award_counter):05d}_7526"
        writer.writerow(full)
    return "﻿".encode() + out.getvalue().encode("utf-8")


def make_awards_zip(rows: list[dict], *, members: int = 1) -> bytes:
    """The zip USAspending serves: the award summary plus a subawards file we ignore. A large
    download is split across numbered members, each with its own header; ``members`` spreads
    the rows over that many, in order."""
    import zipfile

    buffer = io.BytesIO()
    size = -(-len(rows) // members) if rows else 0
    chunks = [rows[i * size : (i + 1) * size] for i in range(members)] if size else [rows]
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, chunk in enumerate(chunks, 1):
            archive.writestr(
                f"Contracts_PrimeAwardSummaries_2026-09-07_H14M26S52_{index}.csv",
                make_awards_csv(chunk),
            )
        archive.writestr("Contracts_Subawards_2026-09-07_H14M29S15_1.csv", "ignored\r\n")
    return buffer.getvalue()


@pytest.fixture
def seed_awards(
    conn: sqlite3.Connection,
    settings: Settings,
    seed: Callable[[dict | None], None],
    tmp_path: Path,
) -> Callable[[list[dict] | None], AwardsResult]:
    """Seed the notices, then ingest an award file: by default one incumbent contract for the
    fixture's HRSA notice and one older award at the same office by another vendor."""
    from orrery.ingest.awards import ingest_awards

    hrsa = SEARCH_FIXTURE["opportunitiesData"][0]
    default_rows = [
        {
            "solicitation_identifier": hrsa["solicitationNumber"],
            "award_id_piid": "75R60222F00009",
            "award_base_action_date": "2022-09-30",
            "award_latest_action_date": "2025-08-01",
        },
        {
            "recipient_uei": "PHZDZ8SJ5CM1",
            "recipient_name": "CDW GOVERNMENT LLC",
            "cage_code": "1KH72",
            "award_id_piid": "75R60224F00021",
            "solicitation_identifier": "",
            "award_base_action_date": "2024-01-15",
            "award_latest_action_date": "2024-01-15",
            "current_total_value_of_award": "48000.00",
            "type_of_set_aside_code": "",
        },
    ]

    def _seed(rows: list[dict] | None = None) -> AwardsResult:
        seed()
        path = tmp_path / "awards.csv"
        path.write_bytes(make_awards_csv(default_rows if rows is None else rows))
        return ingest_awards(conn, settings, path)

    return _seed


EXTRACT_SAMPLE = Path(__file__).with_name("fixtures") / "sam_entity_extract_sample.txt"


@pytest.fixture
def seed_registrations(
    conn: sqlite3.Connection, settings: Settings, seed_awards: Callable[..., AwardsResult]
) -> Callable[[], EntitiesResult]:
    """Seed notices and awards, then the fixture extract: registrations for the two vendors."""
    from orrery.ingest.entities import ingest_extract

    def _seed() -> EntitiesResult:
        seed_awards()
        return ingest_extract(conn, settings, EXTRACT_SAMPLE)

    return _seed
